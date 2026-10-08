"""
Script to generate a realistic Threat Intelligence DOCX report
for testing the /ingest/file API endpoint via Postman.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

doc = Document()

# ─── Document Title ─────────────────────────────────────────────────────────
title = doc.add_heading("THREAT INTELLIGENCE BRIEFING", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.runs[0]
run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
run.font.size = Pt(20)

sub = doc.add_heading("Classified: Lazarus Group Campaign — Operation GhostPay", level=1)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.color.rgb = RGBColor(0x1F, 0x39, 0x7D)

doc.add_paragraph("Date: October 08, 2026 | Classification: RESTRICTED | Source: Joint Threat Analysis Unit (JTAU)")
doc.add_paragraph("Reference ID: JTAU-2026-LZG-0084")
doc.add_horizontal_line = doc.add_paragraph("─" * 80)

# ─── BLUF Section ───────────────────────────────────────────────────────────
doc.add_heading("1. BOTTOM LINE UP FRONT (BLUF)", level=2)
bluf = doc.add_paragraph(
    "Threat Actor: Lazarus Group (DPRK-linked APT38 subgroup) has launched a multi-stage cyber intrusion "
    "campaign targeting the Financial Core Payment Gateway infrastructure across Southeast Asian financial institutions. "
    "The campaign, codenamed Operation GhostPay, combines spear-phishing delivery, credential theft, "
    "and ransomware payload deployment. Intelligence from SIEM, cyber sensors, and OSINT confirms "
    "active Command and Control (C2) communication from malicious IP 203.0.113.199 and domain evil-command-node.xyz. "
    "Immediate defensive action and threat cluster escalation is recommended."
)
bluf.style.font.size = Pt(11)

# ─── Threat Actor ───────────────────────────────────────────────────────────
doc.add_heading("2. THREAT ACTOR PROFILE", level=2)
actor_data = [
    ("Name", "Lazarus Group / APT38"),
    ("Attribution", "Democratic People's Republic of Korea (DPRK)"),
    ("Motivation", "Financial gain, espionage, disruption of Western banking systems"),
    ("TTP Framework", "MITRE ATT&CK: T1566 (Phishing), T1055 (Process Injection), T1486 (Data Encrypted for Impact)"),
    ("Known Campaigns", "Operation AppleJeus (2018), WannaCry (2017), SWIFT banking heist (2016)"),
    ("Current Campaign", "Operation GhostPay — targeting financial payment gateways (Oct 2026)"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Field"
hdr[1].text = "Details"
for field, detail in actor_data:
    row = table.add_row().cells
    row[0].text = field
    row[1].text = detail

doc.add_paragraph("")

# ─── Technical Indicators ───────────────────────────────────────────────────
doc.add_heading("3. INDICATORS OF COMPROMISE (IOCs)", level=2)
doc.add_paragraph(
    "The following indicators have been confirmed across SIEM logs, network sensor telemetry, "
    "and third-party threat intelligence feeds:"
)

ioc_table = doc.add_table(rows=1, cols=3)
ioc_table.style = "Table Grid"
hdr2 = ioc_table.rows[0].cells
hdr2[0].text = "IOC Type"
hdr2[1].text = "Value"
hdr2[2].text = "Confidence"

iocs = [
    ("IP Address", "203.0.113.199", "98%"),
    ("IP Address", "198.51.100.45", "91%"),
    ("Domain", "evil-command-node.xyz", "96%"),
    ("Domain", "ghostpay-loader.net", "88%"),
    ("File Hash (MD5)", "e99a18c428cb38d5f260853678922e03", "99%"),
    ("File Hash (SHA256)", "3f5a2c7b8e1d4f09ac6b2e8f5c0d3a17b2e9f6c4d8a1e3b7f0c5a9d2e8b4f1a3", "99%"),
    ("File Name", "payment_gateway_update.exe", "85%"),
    ("Threat Actor", "Lazarus Group", "95%"),
    ("Target System", "Financial Core Payment Gateway Server", "100%"),
]
for ioc_type, value, conf in iocs:
    row = ioc_table.add_row().cells
    row[0].text = ioc_type
    row[1].text = value
    row[2].text = conf

doc.add_paragraph("")

# ─── Attack Timeline ─────────────────────────────────────────────────────────
doc.add_heading("4. ATTACK TIMELINE", level=2)
timeline_data = [
    ("Day 1 - 00:00", "Spear-phishing email with malicious attachment sent to 3 financial institution employees"),
    ("Day 1 - 08:30", "Victim opens attachment; Cobalt Strike stager executes in memory"),
    ("Day 1 - 09:00", "C2 beacon established to 203.0.113.199 over HTTPS port 443"),
    ("Day 2 - 14:00", "Lateral movement initiated; credential harvesting via Mimikatz"),
    ("Day 3 - 02:15", "Ransomware payload (payment_gateway_update.exe) deployed on PAY-GW-01"),
    ("Day 3 - 02:20", "Mass file encryption initiated on Financial Core Payment Gateway Server"),
    ("Day 3 - 02:25", "Data exfiltration detected to ghostpay-loader.net (198.51.100.45)"),
    ("Day 3 - 03:00", "SIEM alert triggered; automated correlation engine activates"),
    ("Day 3 - 03:05", "Threat cluster CRITICAL escalation; BLUF report auto-generated"),
]
for time, event in timeline_data:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{time}: ").bold = True
    p.add_run(event)

doc.add_paragraph("")

# ─── Affected Systems ────────────────────────────────────────────────────────
doc.add_heading("5. AFFECTED TARGETS", level=2)
doc.add_paragraph(
    "Primary Target: Financial Core Payment Gateway Server (Host: PAY-GW-01)\n"
    "Secondary Target: Internal AD Domain Controller (Host: DC-CORP-01)\n"
    "Tertiary Target: Customer Database Server (Host: DB-CUST-02)\n"
    "Sector: Banking and Financial Services\n"
    "Geographic Region: Southeast Asia — Singapore, Indonesia, Vietnam"
)

# ─── Recommendations ─────────────────────────────────────────────────────────
doc.add_heading("6. RECOMMENDED ACTIONS", level=2)
recommendations = [
    "Immediately block IP 203.0.113.199 and 198.51.100.45 on all edge firewalls and perimeter WAFs.",
    "Sinkhole DNS for domains: evil-command-node.xyz and ghostpay-loader.net.",
    "Scan all systems for file hash e99a18c428cb38d5f260853678922e03 (MD5) and isolate affected hosts.",
    "Reset all privileged account credentials and enable hardware MFA enforcement.",
    "Deploy EDR signatures for Cobalt Strike Beacon process injection patterns (T1055).",
    "Enable enhanced SIEM alerting for SWIFT transaction anomalies and abnormal fund transfers.",
    "Notify relevant national CERT and Financial Intelligence Units within 24 hours.",
    "Human analyst review of auto-generated BLUF report is MANDATORY before cluster escalation closure.",
]
for r in recommendations:
    doc.add_paragraph(r, style="List Bullet")

doc.add_paragraph("")

# ─── Analyst Notes ───────────────────────────────────────────────────────────
doc.add_heading("7. ANALYST NOTES", level=2)
doc.add_paragraph(
    "This briefing was machine-generated by the Sarvatobhadra Threat Intelligence Correlation Platform "
    "using multi-source evidence correlation (SIEM: 4 alerts, Cyber Sensor: 2 events, OSINT: 3 feeds, Intel Report: 1 document). "
    "HUMAN-IN-THE-LOOP REVIEW IS REQUIRED before operational response escalation. "
    "The correlation engine assigned a Threat Score of 94/100 (CRITICAL priority) based on multi-source evidence "
    "and Lazarus Group TTPs historical pattern matching."
)

# ─── Footer ──────────────────────────────────────────────────────────────────
doc.add_paragraph("─" * 80)
footer_para = doc.add_paragraph("Sarvatobhadra Threat Intelligence Correlation Platform | Auto-Generated BLUF | Classification: RESTRICTED")
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ─── Save ────────────────────────────────────────────────────────────────────
output_path = r"src\main\resources\sample-data\LazarusGroup_GhostPay_Intel_Report.docx"
doc.save(output_path)
print(f"DOCX report saved to: {output_path}")
